"""
Test cases for GitHub PR Tool.
"""

import json
from pathlib import Path

from openpyxl import load_workbook


def test_prs_json_file_exists() -> None:
    """
    Verify prs.json exists.
    """

    assert Path("prs.json").exists()


def test_pr_details_json_file_exists() -> None:
    """
    Verify pr_details.json exists.
    """

    assert Path("pr_details.json").exists()


def test_prs_excel_file_exists() -> None:
    """
    Verify prs.xlsx exists.
    """

    assert Path("prs.xlsx").exists()


def test_pr_details_excel_file_exists() -> None:
    """
    Verify pr_details.xlsx exists.
    """

    assert Path("pr_details.xlsx").exists()


def test_prs_json_contains_required_keys() -> None:
    """
    Verify prs.json contains
    required fields.
    """

    with open(
        "prs.json",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    assert len(data) > 0

    assert "pr_number" in data[0]
    assert "title" in data[0]
    assert "author" in data[0]


def test_pr_details_json_contains_required_keys() -> None:
    """
    Verify pr_details.json contains
    required fields.
    """

    with open(
        "pr_details.json",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    assert len(data) > 0

    assert "pr_number" in data[0]
    assert "title" in data[0]
    assert "description" in data[0]
    assert "author" in data[0]
    assert "files_changed" in data[0]
    assert "files_added" in data[0]
    assert "files_deleted" in data[0]


def test_prs_excel_headers() -> None:
    """
    Verify prs.xlsx headers.
    """

    workbook = load_workbook(
        "prs.xlsx"
    )

    worksheet = workbook.active

    assert worksheet["A1"].value == "PR Number"
    assert worksheet["B1"].value == "Title"
    assert worksheet["C1"].value == "Author"


def test_pr_details_excel_headers() -> None:
    """
    Verify pr_details.xlsx headers.
    """

    workbook = load_workbook(
        "pr_details.xlsx"
    )

    worksheet = workbook.active

    assert worksheet["A1"].value == "PR Number"
    assert worksheet["B1"].value == "Title"
    assert worksheet["C1"].value == "Description"
    assert worksheet["D1"].value == "Author"
    assert worksheet["E1"].value == "Files Changed"
    assert worksheet["F1"].value == "Files Added"
    assert worksheet["G1"].value == "Files Deleted"
